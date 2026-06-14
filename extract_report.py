from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("/Users/leonid.ivanenkov/Downloads/AVALAB, Fastboard, Avalin, Itrielt _ ноябрь СМИ_02-12-2025.xlsx")
OUT = Path("report_data.json")


def clean(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        return value.replace("\xa0", " ").strip()
    return value


def rows(ws):
    return [[clean(cell.value) for cell in row] for row in ws.iter_rows()]


def nonempty(values):
    return [row for row in values if any(v not in (None, "") for v in row)]


def parse_content(wb):
    ws = wb["Содержание"]
    result = {}
    for row in ws.iter_rows(values_only=True):
        key = clean(row[1]) if len(row) > 1 else None
        value = clean(row[2]) if len(row) > 2 else None
        if key and value is not None:
            result[key.rstrip(":")] = value
    return result


def parse_summary(wb):
    ws = wb["Сводные данные"]
    result = []
    for row in ws.iter_rows(values_only=True):
        brand = clean(row[1])
        if not brand or brand in {"Сводные данные"}:
            continue
        messages = row[2] or 0
        media_index = row[3] or 0
        negative = row[4] or 0
        neutral = row[6] or 0
        positive = row[8] or 0
        main_role = row[10] or 0
        citations = row[12] or 0 if len(row) > 12 else 0
        result.append(
            {
                "brand": brand,
                "messages": messages,
                "media_index": media_index,
                "negative": negative,
                "neutral": neutral,
                "positive": positive,
                "negative_share": negative / messages if messages else 0,
                "neutral_share": neutral / messages if messages else 0,
                "positive_share": positive / messages if messages else 0,
                "main_role": main_role,
                "main_role_share": main_role / messages if messages else 0,
                "citations": citations,
                "citation_share": citations / messages if messages else 0,
                "media_index_per_mention": media_index / messages if messages else 0,
            }
        )
    return result


def parse_daily_matrix(wb, sheet_name):
    ws = wb[sheet_name]
    header_row = None
    for row in ws.iter_rows(values_only=True):
        vals = [clean(v) for v in row]
        if len(vals) > 2 and vals[1] == "Дата":
            header_row = vals
            break
    if not header_row:
        return []
    headers = header_row[1:7]
    result = []
    for row in ws.iter_rows(values_only=True):
        date = clean(row[1])
        if not date or date == "Дата" or not str(date).startswith("202"):
            continue
        item = {"date": date}
        for idx, name in enumerate(headers[1:], start=2):
            item[name] = row[idx] or 0
        result.append(item)
    return result


def parse_table(wb, sheet_name, min_row, max_col=None):
    ws = wb[sheet_name]
    out = []
    started = False
    for row in ws.iter_rows(values_only=True):
        vals = [clean(v) for v in row[:max_col]] if max_col else [clean(v) for v in row]
        vals = vals[1:] if vals and vals[0] is None else vals
        if not started:
            meaningful = [v for v in vals if v not in (None, "")]
            if len(meaningful) >= 2:
                started = True
            else:
                continue
        if any(v not in (None, "") for v in vals):
            out.append(vals)
    return out


def table_to_records(table):
    if not table:
        return []
    header = table[0]
    records = []
    for row in table[1:]:
        if not any(v not in (None, "") for v in row):
            continue
        records.append({str(header[i]): row[i] if i < len(row) else None for i in range(len(header))})
    return records


def main():
    wb = load_workbook(SOURCE, data_only=True)

    top_events = table_to_records(parse_table(wb, "Инфоповоды", 3, 6))
    top_sources_count = table_to_records(parse_table(wb, "СМИ по количеству", 3, 7))
    top_sources_index = table_to_records(parse_table(wb, "СМИ по МедиаИндексу", 3, 7))
    regions_raw = parse_table(wb, "Регионы", 3, 12)
    words = table_to_records(parse_table(wb, "Слова", 3, 7))
    authors = table_to_records(parse_table(wb, "Авторы", 3, 8))
    industries = table_to_records(parse_table(wb, "СМИ по отраслям", 3, 7))
    rubrics = table_to_records(parse_table(wb, "Рубрики", 3, 7))
    genres = table_to_records(parse_table(wb, "Жанры", 3, 7))

    data = {
        "content": parse_content(wb),
        "summary": parse_summary(wb),
        "daily_mentions": parse_daily_matrix(wb, "Динамика"),
        "daily_media_index": parse_daily_matrix(wb, "МедиаИндекс"),
        "daily_main_role": parse_daily_matrix(wb, "Главная роль"),
        "daily_citation": parse_daily_matrix(wb, "Цитирование"),
        "top_events": top_events[:10],
        "top_sources_count": top_sources_count[:20],
        "top_sources_index": top_sources_index[:20],
        "regions_raw": regions_raw,
        "words": words[:30],
        "authors": authors[:20],
        "industries": industries,
        "rubrics": rubrics,
        "genres": genres,
    }

    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
