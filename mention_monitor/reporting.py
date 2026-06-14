from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from .db import dashboard_stats, latest_mentions


HEADER = "1F4E5F"
LIGHT = "E9F1F4"
TEXT = "172033"
ACCENT = "21C6D3"


def _source_bucket(source: str, url: str | None = None) -> str:
    lower = f"{source or ''} {url or ''}".lower()
    if any(x in lower for x in ["vk.com", "telegram", "t.me", "youtube", "instagram", "dzen.ru", "zen.yandex"]):
        return "Соцплатформы"
    sports_markers = [
        "championat.com",
        "sports.ru",
        "sport24.ru",
        "sport-express.ru",
        "sportbox.ru",
        "matchtv.ru",
        "rsport.ria.ru",
        "stavkinasport",
        "ставкинаспорт",
        "небесная грация",
        "вфв",
        "всероссийская федерация волейбола",
    ]
    if any(x in lower for x in sports_markers):
        return "Спортивные медиа"
    if any(x in lower for x in ["ria.ru", "риа новости", "tass", "тасс", "interfax", "интерфакс", "rbc", "рбк", "lenta.ru", "лента.ру", "kommersant", "коммерсант", "vedomosti", "ведомости", "rg.ru", "российская газета", "iz.ru", "известия", "forbes", "форбс"]):
        return "Федеральные СМИ"
    if any(x in lower for x in ["клопс", "klops", "калининград", "kgd", "baltic", "область", "регион", "nia", "newkaliningrad", "новый калининград", "obltv", "orenburg", "нижегород", "тюмен"]):
        return "Региональные СМИ"
    if any(x in lower for x in ["google news", "news.google", "mail", "dzen", "яндекс", "yandex"]):
        return "Агрегаторы"
    if any(x in lower for x in ["2gis", "reputation", "companium", "fbc.ru", "cataloxy", "yp.ru", "wikipedia"]):
        return "Справочники"
    if any(x in lower for x in [".ru", ".рф", "news", "медиа", "газета"]):
        return "Онлайн-СМИ"
    return "Другие площадки"


def _safe_int(row, key: str) -> int:
    try:
        return int(row[key] or 0) if key in row.keys() else 0
    except (TypeError, ValueError):
        return 0


def _style_header(ws, row: int, columns: int) -> None:
    for cell in ws[row][:columns]:
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _add_table(ws, name: str) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def export_xlsx(conn, output_dir: str = "outputs/reports", **filters) -> str:
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    filename = f"mentions_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    full_path = output_path / filename

    stats = dashboard_stats(conn, **filters)
    rows = latest_mentions(conn, limit=5000, **filters)
    total_engagement = sum(_safe_int(row, "likes") + _safe_int(row, "reposts") + _safe_int(row, "comments") + _safe_int(row, "views") for row in rows)
    measured_social = sum(1 for row in rows if _safe_int(row, "likes") + _safe_int(row, "reposts") + _safe_int(row, "comments") + _safe_int(row, "views") > 0)

    wb = Workbook()
    ws = wb.active
    ws.title = "Дашборд"
    ws["A1"] = "Отчет по упоминаниям"
    ws["A1"].font = Font(size=18, bold=True, color=TEXT)
    ws["A2"] = "Фильтр"
    ws["B2"] = filters.get("project") or "все проекты"
    ws["A3"] = "Всего"
    ws["B3"] = stats["total"]
    ws["A4"] = "Позитив"
    ws["B4"] = stats["by_sentiment"].get("positive", 0)
    ws["A5"] = "Нейтрально"
    ws["B5"] = stats["by_sentiment"].get("neutral", 0)
    ws["A6"] = "Негатив"
    ws["B6"] = stats["by_sentiment"].get("negative", 0)
    ws["A7"] = "Уникальные источники"
    ws["B7"] = stats.get("unique_sources", 0)
    ws["A8"] = "Вовлеченность"
    ws["B8"] = total_engagement
    ws["A9"] = "Публикаций с соцметриками"
    ws["B9"] = measured_social
    for cell in ws["A2:A9"]:
        cell[0].fill = PatternFill("solid", fgColor=LIGHT)
        cell[0].font = Font(bold=True, color=TEXT)

    ws["D2"] = "Динамика"
    ws["D2"].font = Font(bold=True, color=TEXT)
    ws.append([])
    start = 12
    ws.cell(start, 4, "Дата")
    ws.cell(start, 5, "Упоминания")
    for idx, row in enumerate(stats["daily"], start=start + 1):
        ws.cell(idx, 4, row["day"])
        ws.cell(idx, 5, row["n"])
    _style_header(ws, start, 5)
    if stats["daily"]:
        chart = LineChart()
        chart.title = "Динамика упоминаний"
        chart.y_axis.title = "Кол-во"
        chart.x_axis.title = "Дата"
        data = Reference(ws, min_col=5, min_row=start, max_row=start + len(stats["daily"]))
        cats = Reference(ws, min_col=4, min_row=start + 1, max_row=start + len(stats["daily"]))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "G3")

    media_counts: dict[str, dict[str, int]] = {}
    media_engagement: dict[str, int] = {}
    source_counts: dict[str, int] = {}
    source_engagement: dict[str, int] = {}
    for item in rows:
        source = item["source"] or "unknown"
        bucket = _source_bucket(source, item["url"])
        media_counts.setdefault(bucket, {})
        media_counts[bucket][source] = media_counts[bucket].get(source, 0) + 1
        source_counts[source] = source_counts.get(source, 0) + 1
        engagement = _safe_int(item, "likes") + _safe_int(item, "reposts") + _safe_int(item, "comments") + _safe_int(item, "views")
        media_engagement[bucket] = media_engagement.get(bucket, 0) + engagement
        source_engagement[source] = source_engagement.get(source, 0) + engagement

    media = wb.create_sheet("Карта медиа-поля")
    media.append(["Тип площадки", "Публикации", "Источники", "Вовлеченность"])
    for bucket, sources in sorted(media_counts.items(), key=lambda item: -sum(item[1].values())):
        media.append([bucket, sum(sources.values()), len(sources), media_engagement.get(bucket, 0)])
    _style_header(media, 1, 4)
    _add_table(media, "MediaMap")
    if media.max_row > 1:
        chart = BarChart()
        chart.title = "Публикации по типам площадок"
        chart.y_axis.title = "Публикации"
        data = Reference(media, min_col=2, min_row=1, max_row=media.max_row)
        cats = Reference(media, min_col=1, min_row=2, max_row=media.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        media.add_chart(chart, "F2")

    src = wb.create_sheet("Источники")
    src.append(["Тип площадки", "Источник", "Упоминания", "Вовлеченность"])
    for source, count in sorted(source_counts.items(), key=lambda item: (-item[1], item[0].lower())):
        src.append([_source_bucket(source), source, count, source_engagement.get(source, 0)])
    _style_header(src, 1, 4)
    _add_table(src, "Sources")
    if source_counts:
        chart = BarChart()
        chart.title = "Топ источников"
        max_chart_row = min(src.max_row, 16)
        data = Reference(src, min_col=3, min_row=1, max_row=max_chart_row)
        cats = Reference(src, min_col=2, min_row=2, max_row=max_chart_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        src.add_chart(chart, "F2")

    data_ws = wb.create_sheet("Упоминания")
    headers = [
        "Дата сбора",
        "Дата публикации",
        "Проект",
        "Запрос",
        "Тональность",
        "Тип площадки",
        "Источник",
        "Лайки",
        "Репосты",
        "Комментарии",
        "Просмотры",
        "Вовлеченность",
        "Заголовок",
        "Фрагмент",
        "Ссылка",
        "Язык",
    ]
    data_ws.append(headers)
    for item in rows:
        data_ws.append(
            [
                item["collected_at"],
                item["published_at"],
                item["project"],
                item["query"],
                item["sentiment"],
                _source_bucket(item["source"], item["url"]),
                item["source"],
                _safe_int(item, "likes"),
                _safe_int(item, "reposts"),
                _safe_int(item, "comments"),
                _safe_int(item, "views"),
                _safe_int(item, "likes") + _safe_int(item, "reposts") + _safe_int(item, "comments") + _safe_int(item, "views"),
                item["title"],
                item["snippet"],
                item["url"],
                item["language"],
            ]
        )
    _style_header(data_ws, 1, len(headers))
    _add_table(data_ws, "Mentions")

    widths = {
        "A": 22,
        "B": 22,
        "C": 28,
        "D": 28,
        "E": 14,
        "F": 20,
        "G": 28,
        "H": 12,
        "I": 12,
        "J": 14,
        "K": 12,
        "L": 16,
        "M": 48,
        "N": 60,
        "O": 70,
        "P": 12,
    }
    for sheet in wb.worksheets:
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            sheet.column_dimensions[letter].width = widths.get(letter, 18)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    wb.save(full_path)
    return str(full_path)
