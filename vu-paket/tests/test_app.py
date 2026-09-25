import dataclasses
import io
import json
import threading
import unittest
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from datetime import date
from http.server import ThreadingHTTPServer

import docx

from vupaket.config import load_settings
from vupaket.db import Store
from vupaket.documents import SAMPLE_DATA, build_documents, build_ics, initials, plan_rows
from vupaket.payments import PaymentError, apply_payment, process_webhook
from vupaket.render import build_zip, to_docx
from vupaket.server import RateLimiter, make_handler
from vupaket.validation import inn_is_valid, validate

FORM = {
    "org_full": 'Общество с ограниченной ответственностью "Ромашка"',
    "org_short": 'ООО "Ромашка"',
    "inn": "7701000001",
    "city": "Москва",
    "head_position": "Генеральный директор",
    "head_name": "Смирнов Олег Петрович",
    "mode": "appoint",
    "resp_position": "Бухгалтер",
    "resp_name": "Петрова Анна Сергеевна",
    "staff_total": "12",
    "commissariat": "Военный комиссариат Центрального района г. Москвы",
    "year": str(date.today().year + 1),
    "sverka_month": "3",
    "submit_method": "epgu",
    "email": "buh@romashka.ru",
    "consent": "yes",
}


class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *args, **kwargs):
        return None


OPENER = urllib.request.build_opener(NoRedirect)


class ValidationTests(unittest.TestCase):
    def test_inn_checksum(self):
        self.assertTrue(inn_is_valid("7701000001"))
        self.assertFalse(inn_is_valid("7701000002"))
        self.assertFalse(inn_is_valid("12345"))

    def test_valid_form_and_quotes(self):
        data, errors = validate(FORM)
        self.assertEqual(errors, {})
        self.assertEqual(data["org_short"], "ООО «Ромашка»")
        self.assertEqual(data["staff_total"], 12)

    def test_errors(self):
        bad = {**FORM, "inn": "7701000002", "staff_total": "600", "consent": "", "resp_name": "", "head_name": "иванов"}
        _, errors = validate(bad)
        for f in ("inn", "staff_total", "consent", "resp_name", "head_name"):
            self.assertIn(f, errors)

    def test_self_mode_needs_no_responsible(self):
        _, errors = validate({**FORM, "mode": "self", "resp_name": "", "resp_position": ""})
        self.assertEqual(errors, {})


class DocumentTests(unittest.TestCase):
    def test_initials(self):
        self.assertEqual(initials("Иванов Иван Иванович"), "И.И. Иванов")

    def test_docx_contains_company_and_five_day_deadline(self):
        docs = build_documents(SAMPLE_DATA)
        self.assertEqual(len(docs), 6)
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(to_docx(docs[0]))).paragraphs)
        self.assertIn("ООО «Пример»", text)
        self.assertIn("в течение 5 дней", text)
        self.assertIn("Петрова Анна Сергеевна", text)

    def test_self_mode_order(self):
        d = {**SAMPLE_DATA, "mode": "self", "deputy_name": "", "deputy_position": ""}
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(to_docx(build_documents(d)[0]))).paragraphs)
        self.assertIn("оставляю за собой", text)
        self.assertNotIn("Петрова", text)

    def test_plan_dates(self):
        rows = plan_rows(SAMPLE_DATA)
        self.assertEqual(len(rows), 13)
        self.assertIn("март 2027", rows[7][1])
        self.assertIn("До 31 декабря 2027", rows[-1][1])

    def test_ics_lines_are_folded(self):
        ics = build_ics(SAMPLE_DATA, "x")
        for line in ics.split(b"\r\n"):
            self.assertLessEqual(len(line), 75)
        self.assertIn(b"RRULE:FREQ=YEARLY", ics)

    def test_zip(self):
        zf = zipfile.ZipFile(io.BytesIO(build_zip(SAMPLE_DATA, "1")))
        self.assertEqual(len(zf.namelist()), 7)


class ServerCase(unittest.TestCase):
    payments = False

    def setUp(self):
        base = load_settings()
        self.settings = dataclasses.replace(
            base, db_path=":memory:", demo_payments=not self.payments, admin_key="k",
            yookassa_shop_id="shop" if self.payments else "", yookassa_secret_key="sk" if self.payments else "",
            public_url="http://test",
        )
        self.store = Store(":memory:")
        self.remote = {}  # имитация ЮKassa: payment_id -> payment

        def fake_fetch(s, pid):
            if pid not in self.remote:
                raise PaymentError("not found")
            return self.remote[pid]

        def fake_create(s, order, return_url):
            pid = f"p{len(self.remote) + 1}"
            self.remote[pid] = {"id": pid, "status": "pending", "paid": False,
                                "amount": {"value": f"{order['price_rub']:.2f}", "currency": "RUB"},
                                "metadata": {"order_token": order["token"]},
                                "confirmation": {"confirmation_url": f"https://yoomoney.ru/checkout/{pid}"}}
            return pid, f"https://yoomoney.ru/checkout/{pid}"

        self.limiter = RateLimiter()
        handler = make_handler(self.settings, self.store, self.limiter, fetch=fake_fetch, create=fake_create)
        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), handler)
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()
        self.base = f"http://127.0.0.1:{self.httpd.server_address[1]}"

    def tearDown(self):
        self.httpd.shutdown()
        self.httpd.server_close()

    def req(self, path, data=None, headers=None):
        body = urllib.parse.urlencode(data).encode() if isinstance(data, dict) else data
        r = urllib.request.Request(self.base + path, data=body, headers=headers or {})
        try:
            resp = OPENER.open(r)
            return resp.status, dict(resp.headers), resp.read()
        except urllib.error.HTTPError as exc:
            return exc.code, dict(exc.headers), exc.read()

    def create_order(self):
        status, headers, _ = self.req("/order", FORM, {"Cookie": "vu_utm=utm_source%3Dyandex"})
        self.assertEqual(status, 303)
        return headers["Location"]


class DemoFlowTests(ServerCase):
    def test_pages(self):
        for path in ("/", "/form", "/sample", "/offer", "/privacy", "/healthz"):
            self.assertEqual(self.req(path)[0], 200, path)
        self.assertEqual(self.req("/nope")[0], 404)
        self.assertEqual(self.req("/order/short")[0], 404)
        self.assertEqual(self.req("/order/" + "a" * 24)[0], 404)

    def test_invalid_form_returns_422(self):
        status, _, body = self.req("/order", {**FORM, "inn": "1"})
        self.assertEqual(status, 422)
        self.assertIn("ИНН", body.decode())

    def test_full_demo_flow(self):
        loc = self.create_order()
        status, _, body = self.req(loc)
        self.assertEqual(status, 200)
        self.assertIn("Демо-режим", body.decode())
        # До оплаты скачать нельзя
        status, headers, _ = self.req(loc + "/download")
        self.assertEqual((status, headers["Location"]), (303, loc))
        self.assertEqual(self.req(loc + "/pay", {})[0], 303)
        status, _, body = self.req(loc)
        self.assertIn("пакет оплачен", body.decode())
        status, headers, data = self.req(loc + "/download")
        self.assertEqual(status, 200)
        self.assertEqual(headers["Content-Type"], "application/zip")
        self.assertEqual(len(zipfile.ZipFile(io.BytesIO(data)).namelist()), 7)
        # Демо-оплата не считается выручкой
        self.assertEqual(self.store.revenue()["n"], 0)
        # Админка по ключу
        self.assertEqual(self.req("/admin?key=bad")[0], 404)
        self.assertEqual(self.req("/admin?key=k")[0], 200)

    def test_edit_order(self):
        loc = self.create_order()
        status, headers, _ = self.req(loc + "/edit", {**FORM, "city": "Казань"})
        self.assertEqual(status, 303)
        order = self.store.get_order(loc.split("/")[-1])
        self.assertEqual(order["data"]["city"], "Казань")

    def test_analytics_has_no_personal_data(self):
        loc = self.create_order()
        self.req(loc)
        self.req(loc + "/pay", {})
        dump = json.dumps([dict(r) for r in self.store.conn.execute("SELECT * FROM events")], ensure_ascii=False)
        for secret in ("Ромашка", "7701000001", "buh@romashka.ru", "Смирнов", loc.split("/")[-1]):
            self.assertNotIn(secret, dump)
        self.assertIn("yandex", dump)

    def test_rate_limit(self):
        for _ in range(15):
            self.req("/order", {**FORM, "inn": "1"})
        self.assertEqual(self.req("/order", FORM)[0], 429)


class PaymentFlowTests(ServerCase):
    payments = True

    def test_payment_via_redirect_and_webhook(self):
        loc = self.create_order()
        token = loc.split("/")[-1]
        status, headers, _ = self.req(loc + "/pay", {})
        self.assertEqual(status, 303)
        self.assertTrue(headers["Location"].startswith("https://yoomoney.ru/checkout/"))
        # Повторный клик ведёт на тот же платёж
        self.assertEqual(self.req(loc + "/pay", {})[1]["Location"], headers["Location"])
        self.assertEqual(len(self.remote), 1)
        # Вернулись без оплаты — ждём подтверждения
        self.assertIn("Ждём подтверждение", self.req(loc)[2].decode())
        # Вебхук с поддельным «succeeded» в теле не помогает: статус берём из API
        fake = json.dumps({"object": {"id": "p1", "status": "succeeded"}}).encode()
        self.assertEqual(self.req("/yookassa/webhook", fake, {"Content-Type": "application/json"})[0], 200)
        self.assertEqual(self.store.get_order(token)["status"], "pending")
        # Реальная оплата
        self.remote["p1"].update(status="succeeded", paid=True)
        self.assertEqual(self.req("/yookassa/webhook", fake)[0], 200)
        self.assertEqual(self.store.get_order(token)["status"], "paid")
        self.assertEqual(self.req("/yookassa/webhook", fake)[0], 200)  # повтор — без эффекта
        self.assertEqual(self.store.revenue()["n"], 1)
        self.assertEqual(self.req(loc + "/download")[0], 200)

    def test_paid_on_return_without_webhook(self):
        loc = self.create_order()
        self.req(loc + "/pay", {})
        self.remote["p1"].update(status="succeeded", paid=True)
        self.assertIn("пакет оплачен", self.req(loc)[2].decode())

    def test_canceled_payment_returns_to_preview(self):
        loc = self.create_order()
        self.req(loc + "/pay", {})
        self.remote["p1"]["status"] = "canceled"
        self.assertIn("Проверьте данные", self.req(loc)[2].decode())
        self.req(loc + "/pay", {})
        self.assertEqual(len(self.remote), 2)

    def test_amount_mismatch_is_rejected(self):
        loc = self.create_order()
        token = loc.split("/")[-1]
        payment = {"id": "x", "status": "succeeded", "paid": True, "amount": {"value": "1.00", "currency": "RUB"},
                   "metadata": {"order_token": token}}
        self.assertEqual(apply_payment(self.store, payment), "amount-mismatch")
        self.assertEqual(self.store.get_order(token)["status"], "new")

    def test_webhook_api_failure_asks_retry(self):
        body = json.dumps({"object": {"id": "missing"}}).encode()
        self.assertEqual(self.req("/yookassa/webhook", body)[0], 503)
        self.assertEqual(process_webhook(self.settings, self.store, b"{bad"), "bad-json")



class CreatePaymentPayloadTest(unittest.TestCase):
    def test_payload_has_receipt_and_token(self):
        from vupaket import payments
        captured = {}

        def fake_request(s, method, path, payload=None, idem_key=None):
            captured.update(method=method, path=path, payload=payload, idem=idem_key)
            return {"id": "p1", "confirmation": {"confirmation_url": "https://yoomoney.ru/x"}}

        orig = payments._request
        payments._request = fake_request
        try:
            s = dataclasses.replace(load_settings(), yookassa_receipt=True, yookassa_vat_code=1)
            pid, url = payments.create_payment(s, {"id": 5, "token": "t" * 24, "price_rub": 2490, "email": "a@b.ru"}, "http://r")
        finally:
            payments._request = orig
        p = captured["payload"]
        self.assertEqual((pid, url), ("p1", "https://yoomoney.ru/x"))
        self.assertEqual(p["amount"], {"value": "2490.00", "currency": "RUB"})
        self.assertEqual(p["metadata"], {"order_token": "t" * 24})
        self.assertEqual(p["receipt"]["customer"]["email"], "a@b.ru")
        self.assertEqual(p["receipt"]["items"][0]["amount"]["value"], "2490.00")
        self.assertTrue(captured["idem"])


if __name__ == "__main__":
    unittest.main()
