"""Сборка документов: DOCX (python-docx), HTML-предпросмотр и ZIP-архив пакета."""
from __future__ import annotations

import html
import io
import zipfile

from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .documents import Document, build_documents, build_ics

FONT = "Times New Roman"


def _para(doc, text: str = "", *, bold=False, align=None, size=12, indent=False, italic=False, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.space_before = Pt(0)
    if indent:
        fmt.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = FONT
    return p


def _cell_text(cell, text: str, *, bold=False, size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT


def to_docx(document: Document) -> bytes:
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    section = doc.sections[0]
    wide = any(bl[0] == "table" and len(bl[1]) >= 6 for bl in document.blocks)
    if wide:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(2.5 if not wide else 1.5)
    section.right_margin = Cm(1.2)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    for block in document.blocks:
        kind = block[0]
        if kind == "org_header":
            for line in block[1]:
                _para(doc, line, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
            _para(doc, "", space_after=2)
        elif kind == "title":
            _para(doc, block[1], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=2)
        elif kind == "center":
            _para(doc, block[1], align=WD_ALIGN_PARAGRAPH.CENTER)
        elif kind == "h2":
            _para(doc, block[1], bold=True, space_after=6)
        elif kind == "para":
            _para(doc, block[1], align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True)
        elif kind == "item":
            _para(doc, block[1], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif kind == "plain":
            _para(doc, block[1])
        elif kind == "note":
            _para(doc, block[1], italic=True, size=10)
        elif kind == "spacer":
            _para(doc, "")
        elif kind == "right":
            for line in block[1]:
                _para(doc, line, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0, size=11)
            _para(doc, "", space_after=2)
        elif kind == "row2":
            t = doc.add_table(rows=1, cols=2)
            _cell_text(t.cell(0, 0), block[1], size=12)
            _cell_text(t.cell(0, 1), block[2], size=12, align=WD_ALIGN_PARAGRAPH.RIGHT)
        elif kind == "approval":
            left, right = block[1], block[2]
            t = doc.add_table(rows=1, cols=2)
            _cell_text(t.cell(0, 0), "\n".join(left), size=11)
            _cell_text(t.cell(0, 1), "\n".join(right), size=11)
            _para(doc, "", space_after=2)
        elif kind == "sign":
            t = doc.add_table(rows=1, cols=2)
            _cell_text(t.cell(0, 0), block[1], size=12)
            _cell_text(t.cell(0, 1), f"____________ {block[2]}", size=12, align=WD_ALIGN_PARAGRAPH.RIGHT)
        elif kind == "table":
            headers, rows, widths = block[1], block[2], block[3]
            t = doc.add_table(rows=1 + len(rows), cols=len(headers))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(headers):
                _cell_text(t.cell(0, j), h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            for i, row in enumerate(rows, 1):
                for j, value in enumerate(row):
                    _cell_text(t.cell(i, j), value)
            for j, w in enumerate(widths):
                for row in t.rows:
                    row.cells[j].width = Cm(w)
            _para(doc, "", space_after=2)
        else:  # pragma: no cover - защита от опечаток в documents.py
            raise ValueError(f"unknown block {kind}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_html(document: Document, limit_blocks: int | None = None) -> str:
    """Предпросмотр документа. limit_blocks обрезает текст до оплаты."""
    e = html.escape
    parts = [f'<article class="doc"><div class="doc-name">{e(document.title)}</div>']
    blocks = document.blocks if limit_blocks is None else document.blocks[:limit_blocks]
    for block in blocks:
        kind = block[0]
        if kind == "org_header":
            parts.append('<p class="c b">' + "<br>".join(e(x) for x in block[1]) + "</p>")
        elif kind == "title":
            parts.append(f'<p class="c b big">{e(block[1])}</p>')
        elif kind == "center":
            parts.append(f'<p class="c">{e(block[1])}</p>')
        elif kind == "h2":
            parts.append(f'<p class="b">{e(block[1])}</p>')
        elif kind == "para":
            parts.append(f'<p class="j ind">{e(block[1])}</p>')
        elif kind in ("item", "plain"):
            parts.append(f'<p class="j">{e(block[1])}</p>')
        elif kind == "note":
            parts.append(f'<p class="note">{e(block[1])}</p>')
        elif kind == "spacer":
            parts.append('<p class="sp"></p>')
        elif kind == "right":
            parts.append('<p class="r small">' + "<br>".join(e(x) for x in block[1]) + "</p>")
        elif kind == "row2":
            parts.append(f'<div class="row2"><span>{e(block[1])}</span><span>{e(block[2])}</span></div>')
        elif kind == "approval":
            left = "<br>".join(e(x) for x in block[1])
            right = "<br>".join(e(x) for x in block[2])
            parts.append(f'<div class="row2 small"><span>{left}</span><span>{right}</span></div>')
        elif kind == "sign":
            parts.append(f'<div class="row2"><span>{e(block[1])}</span><span>____________ {e(block[2])}</span></div>')
        elif kind == "table":
            head = "".join(f"<th>{e(h)}</th>" for h in block[1])
            body = "".join("<tr>" + "".join(f"<td>{e(v)}</td>" for v in row) + "</tr>" for row in block[2])
            parts.append(f'<div class="tbl"><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>')
    parts.append("</article>")
    return "".join(parts)


def build_zip(data: dict, uid_seed: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for document in build_documents(data):
            zf.writestr(document.filename, to_docx(document))
        zf.writestr("07_Напоминания_о_сроках.ics", build_ics(data, uid_seed))
    return buf.getvalue()
